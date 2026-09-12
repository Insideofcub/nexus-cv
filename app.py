import os
import streamlit as st
import zipfile
import tempfile
from pypdf import PdfReader
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

st.set_page_config(page_title="NexusCV Free", page_icon="⚡", layout="wide")

st.markdown("<h1 style='text-align: center;'>⚡ NexusCV Free Intelligence Engine</h1>", unsafe_allow_html=True)
st.markdown("<p style='text-align: center; color: #9ca3af;'>Upload your zipped resume folder and job description for instant cloud ranking.</p>", unsafe_allow_html=True)
st.markdown("---")

col1, col2 = st.columns(2)

with col1:
    st.markdown("### 📋 1. Job Description")
    jd_file = st.file_uploader("Upload JD (.txt, .pdf, .docx)", type=["txt", "pdf", "docx"], key="jd_free")
    jd_text = ""
    if jd_file:
        if jd_file.name.endswith(".txt"):
            jd_text = jd_file.read().decode("utf-8")
        elif jd_file.name.endswith(".pdf"):
            reader = PdfReader(jd_file)
            jd_text = "\n".join([p.extract_text() for p in reader.pages if p.extract_text()])
        elif jd_file.name.endswith(".docx"):
            doc = Document(jd_file)
            jd_text = "\n".join([p.text for p in doc.paragraphs])

with col2:
    st.markdown("### 📂 2. Candidate Resumes Folder")
    zip_file = st.file_uploader("Upload Entire Resumes Folder (.zip)", type=["zip"], key="zip_upload")

st.markdown("<br>", unsafe_allow_html=True)
if st.button("🚀 Run Free Ranking"):
    if not jd_text.strip():
        st.error("Please upload a job description.")
    elif not zip_file:
        st.error("Please upload your zipped resume folder.")
    else:
        with st.spinner("Extracting and analyzing resumes from folder..."):
            resume_data = []
            resume_texts = []
            
            with tempfile.TemporaryDirectory() as tmp_dir:
                zip_path = os.path.join(tmp_dir, "uploaded_resumes.zip")
                with open(zip_path, "wb") as f:
                    f.write(zip_file.getbuffer())
                
                with zipfile.ZipFile(zip_path, 'r') as zip_ref:
                    zip_ref.extractall(tmp_dir)
                
                # Scan extracted files recursively
                for root, dirs, files in os.walk(tmp_dir):
                    for file in files:
                        if file.lower().endswith(('.pdf', '.docx')) and not file.startswith('._'):
                            file_path = os.path.join(root, file)
                            text = ""
                            try:
                                if file.lower().endswith('.pdf'):
                                    reader = PdfReader(file_path)
                                    text = "\n".join([p.extract_text() for p in reader.pages if p.extract_text()])
                                elif file.lower().endswith('.docx'):
                                    doc = Document(file_path)
                                    text = "\n".join([p.text for p in doc.paragraphs])
                                
                                if text.strip():
                                    with open(file_path, "rb") as f_obj:
                                        file_bytes = f_obj.read()
                                    resume_data.append({"filename": file, "text": text, "bytes": file_bytes})
                                    resume_texts.append(text)
                            except Exception:
                                pass

            if not resume_texts:
                st.error("No valid PDF or Word resumes found inside the uploaded zip file.")
            else:
                vectorizer = TfidfVectorizer(stop_words='english', max_features=10000)
                tfidf_matrix = vectorizer.fit_transform(resume_texts + [jd_text])
                
                query_vector = tfidf_matrix[-1]
                resume_vectors = tfidf_matrix[:-1]
                
                similarities = cosine_similarity(resume_vectors, query_vector).flatten()
                ranked_indices = similarities.argsort()[::-1]
                
                st.markdown("---")
                st.subheader(f"🏆 Top Ranked Candidates ({len(resume_texts)} Processed)")
                
                for rank, idx in enumerate(ranked_indices, 1):
                    candidate = resume_data[idx]
                    score = round(float(similarities[idx]) * 100, 2)
                    
                    with st.expander(f"Rank #{rank}: {candidate['filename']} — Match Score: {score}%"):
                        st.markdown(f"**Snippet:**\n> {candidate['text'][:400]}...")
                        st.download_button(
                            label="📥 Download Resume",
                            data=candidate["bytes"],
                            file_name=candidate["filename"],
                            mime="application/octet-stream",
                            key=f"free_dl_{rank}_{candidate['filename']}"
                        )
